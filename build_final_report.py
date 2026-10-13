import os
import sys
import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER

sys.stdout.reconfigure(encoding='utf-8')

src_path = r"C:\Users\thanh\.gemini\antigravity\brain\7de41738-3605-4b96-bc46-b99f8446a5a7\.user_uploaded\media_1790862856943.docx"
dst_path = r"c:\Users\thanh\Desktop\klcn\Bao_Cao_Thuc_Tap_Nghe_Nghiep_Nguyen_Huy_Thanh_DEMO_FINAL.docx"
img_dir = r"c:\Users\thanh\Desktop\klcn\report_demo_screenshots"

print(f"Loading original document from: {src_path}")
doc = docx.Document(src_path)

def add_figure(target_p, img_path, caption_text):
    """Inserts an image and caption right before target_p."""
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(10)
    img_p.paragraph_format.space_after = Pt(4)
    run_img = img_p.add_run()
    run_img.add_picture(img_path, width=Cm(14.5))
    
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_before = Pt(0)
    cap_p.paragraph_format.space_after = Pt(12)
    run_cap = cap_p.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.bold = True
    run_cap.font.italic = True
    
    target_p._p.addprevious(img_p._p)
    target_p._p.addprevious(cap_p._p)
    return img_p, cap_p

# ============================================================
# STEP 1: LOCATE ALL ANCHORS IN ADVANCE (BEFORE ANY MODIFICATIONS)
# ============================================================
# 1.1 In Section 2.5.1
cap_p2_8 = None
img_p2_8 = None
target_p_251 = None

for i, p in enumerate(doc.paragraphs):
    if i > 150 and p.text.strip().startswith("Hình 2.8:") and "\t" not in p.text:
        cap_p2_8 = p
        img_p2_8 = doc.paragraphs[i - 1]
        target_p_251 = doc.paragraphs[i + 1]
        break

# 1.2 In Section 2.5.2
arch_p_tarot = None
cap_p2_10 = None
img_p2_10 = None
target_p_252 = None

for i, p in enumerate(doc.paragraphs):
    if i > 150 and p.text.strip().startswith("Hình 2.9:") and "\t" not in p.text:
        arch_p_tarot = p
    elif i > 150 and p.text.strip().startswith("Hình 2.10:") and "\t" not in p.text:
        cap_p2_10 = p
        img_p2_10 = doc.paragraphs[i - 1]
        target_p_252 = doc.paragraphs[i + 1]

print(f"Anchor 2.5.1 cap_p2_8: {cap_p2_8.text}")
print(f"Anchor 2.5.1 target_p_251: {target_p_251.text[:40]}...")
print(f"Anchor 2.5.2 arch_p_tarot: {arch_p_tarot.text}")
print(f"Anchor 2.5.2 cap_p2_10: {cap_p2_10.text}")
print(f"Anchor 2.5.2 target_p_252: {target_p_252.text[:40]}...")

# ============================================================
# STEP 2: INSERT AI-AGENT-PLATFORM FIGURES (HÌNH 2.8 -> 2.12)
# ============================================================
print("\n>>> Inserting ai-agent-platform demo figures...")
platform_figures = [
    (os.path.join(img_dir, "05_platform_chat_ui.png"), "Hình 2.8: Giao diện Chat UI bảng điều khiển trung tâm hệ thống Multi-Agent Platform"),
    (os.path.join(img_dir, "06_hiring_candidates_table.png"), "Hình 2.9: Phân hệ Tuyển dụng (Hiring Agent) phân tích yêu cầu và gọi MCP tool tra cứu ứng viên"),
    (os.path.join(img_dir, "07_hiring_candidate_detail.png"), "Hình 2.10: Phân hệ Tuyển dụng truy vấn chi tiết hồ sơ ứng viên UV001 qua MCP tool"),
    (os.path.join(img_dir, "08_employee_department_demo.png"), "Hình 2.11: Phân hệ Nhân sự (Employee Agent) gọi MCP tool tra cứu cơ cấu tổ chức phòng ban"),
    (os.path.join(img_dir, "09_mcp_flow_terminal_log.png"), "Hình 2.12: Nhật ký vết thực thi (Execution Trace) luồng xử lý Root Agent điều phối Domain Agent và MCP Server"),
]

for img_path, caption in platform_figures:
    print(f"  + {caption}")
    add_figure(target_p_251, img_path, caption)

# Remove old mock image 12 & caption 2.8
img_p2_8._p.getparent().remove(img_p2_8._p)
cap_p2_8._p.getparent().remove(cap_p2_8._p)

# ============================================================
# STEP 3: RENUMBER TAROT ARCHITECTURE & INSERT TAROT FIGURES
# ============================================================
print("\n>>> Inserting tarot-viewer-ai demo figures...")
arch_p_tarot.text = "Hình 2.13: Sơ đồ kiến trúc xử lý và luồng trải bài Tarot trong tarot-viewer-ai"
arch_p_tarot.runs[0].font.name = "Times New Roman"
arch_p_tarot.runs[0].font.bold = True
arch_p_tarot.runs[0].font.italic = True
print(f"  Renumbered: {arch_p_tarot.text}")

tarot_figures = [
    (os.path.join(img_dir, "01_tarot_home.png"), "Hình 2.14: Giao diện khởi đầu không gian trải bài Tarot Sanctuary"),
    (os.path.join(img_dir, "02_tarot_spread.png"), "Hình 2.15: Giao diện tương tác lựa chọn sơ đồ trải bài trong tarot-viewer-ai"),
    (os.path.join(img_dir, "03_tarot_result.png"), "Hình 2.16: Giao diện trực quan hóa lá bài và luận giải từng bước điều phối bởi LangGraph"),
    (os.path.join(img_dir, "04_tarot_synthesis.png"), "Hình 2.17: Giao diện tổng kết luận giải toàn diện (Final Synthesis) và lời khuyên hành động"),
]

for img_path, caption in tarot_figures:
    print(f"  + {caption}")
    add_figure(target_p_252, img_path, caption)

# Remove old placeholder image 14 & caption 2.10
img_p2_10._p.getparent().remove(img_p2_10._p)
cap_p2_10._p.getparent().remove(cap_p2_10._p)

# ============================================================
# STEP 4: UPDATE DANH MỤC HÌNH
# ============================================================
print("\n>>> Updating DANH MỤC HÌNH...")
dm_header = None
dmb_header = None

for p in doc.paragraphs:
    if p.text.strip() == "DANH MỤC HÌNH":
        dm_header = p
    elif p.text.strip() == "DANH MỤC BẢNG":
        dmb_header = p
        break

# Remove existing figure entries
cur = dm_header._p.getnext()
while cur is not None and cur != dmb_header._p:
    nxt = cur.getnext()
    # Check if paragraph has text starting with Hình
    p_temp = docx.text.paragraph.Paragraph(cur, doc)
    if p_temp.text.strip().startswith("Hình"):
        cur.getparent().remove(cur)
    cur = nxt

full_figure_list = [
    ("Hình 1.1: Logo chính thức của Tâm Anh Software Solutions", "2"),
    ("Hình 1.2: Sơ đồ cơ cấu tổ chức bộ máy Công ty Tâm Anh Software Solutions", "4"),
    ("Hình 2.1: Kiến trúc luồng xử lý truy vấn của LangChain Framework", "15"),
    ("Hình 2.2: Mô hình StateGraph và luồng điều phối vòng lặp trong LangGraph", "16"),
    ("Hình 2.3: Mô hình điều phối Supervisor Pattern và Agent Handoff trong hệ đa tác tử", "17"),
    ("Hình 2.4: Quy trình trích xuất và truy vấn dữ liệu theo kiến trúc RAG", "18"),
    ("Hình 2.5: Kiến trúc Model Context Protocol (MCP Host, Client, Server, Tools)", "19"),
    ("Hình 2.6: Sơ đồ kiến trúc tổng thể nền tảng ai-agent-platform", "21"),
    ("Hình 2.7: Quy trình phân tích Intent và điều phối của Root Agent đến Hiring Agent", "22"),
    ("Hình 2.8: Giao diện Chat UI bảng điều khiển trung tâm hệ thống Multi-Agent Platform", "23"),
    ("Hình 2.9: Phân hệ Tuyển dụng (Hiring Agent) phân tích yêu cầu và gọi MCP tool tra cứu ứng viên", "23"),
    ("Hình 2.10: Phân hệ Tuyển dụng truy vấn chi tiết hồ sơ ứng viên UV001 qua MCP tool", "24"),
    ("Hình 2.11: Phân hệ Nhân sự (Employee Agent) gọi MCP tool tra cứu cơ cấu tổ chức phòng ban", "25"),
    ("Hình 2.12: Nhật ký vết thực thi (Execution Trace) luồng xử lý Root Agent điều phối Domain Agent và MCP Server", "26"),
    ("Hình 2.13: Sơ đồ kiến trúc xử lý và luồng trải bài Tarot trong tarot-viewer-ai", "27"),
    ("Hình 2.14: Giao diện khởi đầu không gian trải bài Tarot Sanctuary", "28"),
    ("Hình 2.15: Giao diện tương tác lựa chọn sơ đồ trải bài trong tarot-viewer-ai", "28"),
    ("Hình 2.16: Giao diện trực quan hóa lá bài và luận giải từng bước điều phối bởi LangGraph", "29"),
    ("Hình 2.17: Giao diện tổng kết luận giải toàn diện (Final Synthesis) và lời khuyên hành động", "30"),
]

for title, page_str in full_figure_list:
    item_p = doc.add_paragraph()
    item_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    item_p.paragraph_format.space_before = Pt(0)
    item_p.paragraph_format.space_after = Pt(2)
    item_p.paragraph_format.tab_stops.add_tab_stop(5760720, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    
    r1 = item_p.add_run(title)
    r1.font.name = "Times New Roman"
    
    r_tab = item_p.add_run("\t")
    
    r2 = item_p.add_run(page_str)
    r2.font.name = "Times New Roman"
    
    dmb_header._p.addprevious(item_p._p)

print(f"  Inserted {len(full_figure_list)} entries into DANH MỤC HÌNH.")

# ============================================================
# STEP 5: UPDATE TABLE OF CONTENTS (MỤC LỤC)
# ============================================================
print("\n>>> Updating Table of Contents (Mục lục)...")
toc_updates = {
    "2.5. Chi tiết các dự án xây dựng trong quá trình thực tập": "20",
    "2.5.1. Dự án 1: Nền tảng điều phối đa tác tử (ai-agent-platform)": "20",
    "2.5.2. Dự án 2: Ứng dụng AI luận giải Tarot (tarot-viewer-ai)": "27",
    "CHƯƠNG 3: NHẬN XÉT, ĐÁNH GIÁ THỰC TRẠNG": "31",
    "3.1. Các nhận xét, đánh giá thực trạng của quá trình làm việc": "31",
    "3.1.1. Đánh giá về đơn vị thực tập": "31",
    "3.1.2. Tự đánh giá bản thân và kết quả thu nhận được": "31",
    "3.1.3. Những hạn chế và khó khăn gặp phải": "32",
    "3.2. Các kiến nghị và đề xuất": "33",
    "3.2.1. Kiến nghị đối với Nhà trường (Trường ĐH Công Thương TP.HCM)": "33",
    "3.2.2. Kiến nghị đối với Doanh nghiệp (Tâm Anh Software Solutions)": "33",
    "KẾT LUẬN": "35",
    "TÀI LIỆU THAM KHẢO": "36",
}

for p in doc.paragraphs[:78]:
    for key, new_pg in toc_updates.items():
        if p.text.startswith(key):
            if len(p.runs) >= 3:
                p.runs[-1].text = new_pg
            elif "\t" in p.text:
                parts = p.text.split("\t")
                p.text = f"{parts[0]}\t{new_pg}"
            print(f"  TOC: {key} -> {new_pg}")

# ============================================================
# STEP 6: SAVE TO FINAL OUTPUT
# ============================================================
print(f"\nSaving final report to: {dst_path}")
doc.save(dst_path)
print(">>> SUCCESS: Final report generated successfully!")
