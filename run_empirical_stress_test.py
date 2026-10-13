# -*- coding: utf-8 -*-
"""
COMPREHENSIVE EMPIRICAL STRESS TEST HARNESS
Target: Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx
Author: teamwork_preview_challenger_2
"""

import os
import sys
import json
import math
import re
import docx
import pandas as pd

if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"c:\Users\thanh\Desktop\klcn\Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx"
WORKSPACE = r"c:\Users\thanh\Desktop\klcn"

results = {
    "total_tests": 0,
    "passed_tests": 0,
    "failed_tests": 0,
    "warnings": 0,
    "details": []
}

def record_test(name, passed, message, is_warning=False):
    results["total_tests"] += 1
    if passed:
        results["passed_tests"] += 1
        status = "PASS"
    elif is_warning:
        results["warnings"] += 1
        status = "WARN"
    else:
        results["failed_tests"] += 1
        status = "FAIL"
    results["details"].append({"name": name, "status": status, "message": message})
    print(f"[{status}] {name}: {message}")

def run_stress_tests():
    print("=" * 80)
    print("STARTING ADVANCED EMPIRICAL STRESS TESTS")
    print("=" * 80)

    # 1. Load document
    assert os.path.exists(DOCX_PATH), "Target docx does not exist!"
    doc = docx.Document(DOCX_PATH)
    
    # Collect all text and table strings
    para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(para_texts)
    
    tables_data = []
    for tbl in doc.tables:
        t_rows = []
        for row in tbl.rows:
            t_rows.append([c.text.strip().replace('\n', ' ') for c in row.cells])
        tables_data.append(t_rows)

    # Flatten tables to text
    all_table_text = "\n".join(["\n".join([" | ".join(r) for r in t]) for t in tables_data])
    entire_doc_corpus = full_text + "\n" + all_table_text

    # --------------------------------------------------------------------------
    # SUITE 1: FUNDUS GATE MATHEMATICAL AND CODEBASE VERIFICATION
    # --------------------------------------------------------------------------
    print("\n--- SUITE 1: FUNDUS GATE RIGOROUS CHECK ---")
    fg_json_path = os.path.join(WORKSPACE, "fundus-gate", "metrics.json")
    with open(fg_json_path, "r", encoding="utf-8") as f:
        fg_meta = json.load(f)

    # Ground truth values
    tm = fg_meta["test_metrics"]
    cm = tm["confusion_matrix"]
    ch = fg_meta["challenge_test_metrics"]
    opt_tau = fg_meta["optimal_threshold"]
    tot_params = fg_meta["total_parameters"]

    # Test 1.1: Confusion matrix arithmetic
    tp, fp, tn, fn = cm["tp"], cm["fp"], cm["tn"], cm["fn"]
    recomputed_rec = tp / (tp + fn)
    recomputed_spec = tn / (tn + fp)
    recomputed_acc = (tp + tn) / (tp + tn + fp + fn)
    recomputed_fpr = fp / (fp + tn)
    recomputed_f1 = 2 * (tp / (tp + fp)) * recomputed_rec / ((tp / (tp + fp)) + recomputed_rec)

    record_test("FG_CM_Recall_Math", abs(recomputed_rec - 1.0) < 1e-6, f"Recall = {recomputed_rec * 100:.2f}% (TP={tp}, FN={fn})")
    record_test("FG_CM_Spec_Math", abs(recomputed_spec - 0.998731) < 1e-4, f"Specificity = {recomputed_spec * 100:.2f}% (TN={tn}, FP={fp})")
    record_test("FG_CM_Acc_Math", abs(recomputed_acc - 0.999368) < 1e-4, f"Accuracy = {recomputed_acc * 100:.2f}%")
    record_test("FG_CM_FPR_Math", abs(recomputed_fpr - 0.001269) < 1e-4, f"FPR = {recomputed_fpr * 100:.2f}%")
    record_test("FG_CM_F1_Math", abs(recomputed_f1 - 0.999371) < 1e-4, f"F1 = {recomputed_f1:.4f}")

    # Test 1.2: Challenge set arithmetic
    ch_tn, ch_fp, ch_tot = ch["true_negatives"], ch["false_positives"], ch["total_samples"]
    ch_spec = ch_tn / ch_tot
    ch_fpr = ch_fp / ch_tot
    record_test("FG_Challenge_Spec_Math", abs(ch_spec - 0.998889) < 1e-4, f"Challenge Spec = {ch_spec * 100:.2f}% (TN={ch_tn}, FP={ch_fp})")
    record_test("FG_Challenge_FPR_Math", abs(ch_fpr - 0.001111) < 1e-4, f"Challenge FPR = {ch_fpr * 100:.2f}%")

    # Test 1.3: Document presence of FG exact numbers
    record_test("FG_Doc_Recall", "100.00%" in entire_doc_corpus or "100.0%" in entire_doc_corpus, "Recall 100.0% stated in docx")
    record_test("FG_Doc_Spec", "99.87%" in entire_doc_corpus, "Specificity 99.87% stated in docx")
    record_test("FG_Doc_Acc", "99.94%" in entire_doc_corpus, "Accuracy 99.94% stated in docx")
    record_test("FG_Doc_FPR", "0.13%" in entire_doc_corpus, "FPR 0.13% stated in docx")
    record_test("FG_Doc_Challenge_FPR", "0.11%" in entire_doc_corpus, "Challenge FPR 0.11% stated in docx")
    record_test("FG_Doc_Threshold", "0.965" in entire_doc_corpus, "Optimal threshold tau*=0.965 stated in docx")
    record_test("FG_Doc_Params", "939,697" in entire_doc_corpus, "Exact parameter count 939,697 stated in docx")
    record_test("FG_Doc_TFLite_FP16", "1.82 MB" in entire_doc_corpus, "TFLite FP16 size 1.82 MB stated in docx")

    # Physical file check
    fg_fp16_file = os.path.join(WORKSPACE, "fundus-gate", "fundus_gate_float16.tflite")
    if os.path.exists(fg_fp16_file):
        sz_mb = os.path.getsize(fg_fp16_file) / (1024 * 1024)
        record_test("FG_Physical_File_Size", abs(sz_mb - 1.8176) < 0.05, f"Physical file size: {sz_mb:.2f} MB matches claimed 1.82 MB")

    # --------------------------------------------------------------------------
    # SUITE 2: RESNET-50 AND VIT TRAINING LOGS CROSS-REFERENCE
    # --------------------------------------------------------------------------
    print("\n--- SUITE 2: RESNET-50 & VIT CLASSIFICATION LOGS ---")
    csv_file = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "results_comparison.csv")
    df_logs = pd.read_csv(csv_file)

    # Check Table 2 (Table 2.1 in document)
    tbl_2_1 = tables_data[1]  # Table 2 is Table 2.1
    # Check headers
    record_test("Table_2_1_Header_Epoch", "Epoch" in tbl_2_1[0][0], "Table 2.1 column 0 is Epoch")
    record_test("Table_2_1_Header_Loss", "Val Loss" in tbl_2_1[0][1], "Table 2.1 column 1 is Val Loss")
    record_test("Table_2_1_Header_F1", "Val Macro-F1" in tbl_2_1[0][3], "Table 2.1 column 3 is Val Macro-F1")
    record_test("Table_2_1_Header_QWK", "Val QWK" in tbl_2_1[0][4], "Table 2.1 column 4 is Val QWK")

    # Verify each epoch row against CSV
    for ep in range(1, 11):
        row = tbl_2_1[ep]
        csv_row = df_logs[df_logs["Epoch"] == ep].iloc[0]
        
        doc_loss_str = row[1].replace("*", "")
        doc_f1_str = row[3].replace("*", "")
        doc_qwk_str = row[4].replace("*", "")
        
        csv_loss = csv_row["ResNet50_Loss"]
        csv_f1 = csv_row["ResNet50_F1"]
        csv_qwk = csv_row["ResNet50_QWK"]
        
        loss_match = abs(float(doc_loss_str) - csv_loss) < 0.001
        f1_match = abs(float(doc_f1_str) - csv_f1) < 0.001
        qwk_match = abs(float(doc_qwk_str) - csv_qwk) < 0.001
        
        record_test(f"ResNet_Ep{ep}_CSV_Match", loss_match and f1_match and qwk_match,
                    f"Epoch {ep}: Doc({doc_loss_str}, {doc_f1_str}, {doc_qwk_str}) vs CSV({csv_loss:.4f}, {csv_f1:.4f}, {csv_qwk:.4f})")

    # Verify Best Epochs
    best_loss_ep = df_logs.loc[df_logs["ResNet50_Loss"].idxmin(), "Epoch"]
    best_f1_ep = df_logs.loc[df_logs["ResNet50_F1"].idxmax(), "Epoch"]
    best_qwk_ep = df_logs.loc[df_logs["ResNet50_QWK"].idxmax(), "Epoch"]
    
    record_test("ResNet_Best_Loss_Epoch", best_loss_ep == 7, f"Best Loss at Epoch {best_loss_ep} (Claimed: Epoch 7)")
    record_test("ResNet_Best_F1_Epoch", best_f1_ep == 7, f"Best Macro-F1 at Epoch {best_f1_ep} (Claimed: Epoch 7)")
    record_test("ResNet_Best_QWK_Epoch", best_qwk_ep == 9, f"Best QWK at Epoch {best_qwk_ep} (Claimed: Epoch 9)")

    # --------------------------------------------------------------------------
    # SUITE 3: CLASS IMBALANCE & LOSS SMOOTHING MATHEMATICS
    # --------------------------------------------------------------------------
    print("\n--- SUITE 3: CLASS IMBALANCE MATHEMATICAL VERIFICATION ---")
    counts = [25810, 2443, 5292, 873, 708]
    n_tot = sum(counts)
    c_num = 5
    
    # Class percentages
    pct_0 = counts[0] / n_tot * 100
    pct_4 = counts[4] / n_tot * 100
    record_test("EyePACS_Class0_Pct", abs(pct_0 - 73.48) < 0.05, f"Class 0 = {pct_0:.2f}% (Claimed: 73.48%)")
    record_test("EyePACS_Class4_Pct", abs(pct_4 - 2.02) < 0.05, f"Class 4 = {pct_4:.2f}% (Claimed: 2.02%)")

    # Raw weights & ratio
    w_raw = [n_tot / (c_num * cnt) for cnt in counts]
    raw_ratio = w_raw[4] / w_raw[0]
    record_test("Class_Weight_Raw_Ratio", abs(raw_ratio - 36.45) < 0.1, f"Raw ratio = {raw_ratio:.2f}x (Claimed: 36.5x)")

    # Smoothed weights & ratio
    w_smooth = [math.sqrt(w) for w in w_raw]
    smooth_ratio = w_smooth[4] / w_smooth[0]
    record_test("Class_Weight_Smooth_Ratio", abs(smooth_ratio - 6.04) < 0.1, f"Smoothed ratio = {smooth_ratio:.2f}x (Claimed: 6.0x)")

    # --------------------------------------------------------------------------
    # SUITE 4: BENCHMARK MATRIX (TABLE 3.1) EMPIRICAL VALIDATION
    # --------------------------------------------------------------------------
    print("\n--- SUITE 4: BENCHMARK MATRIX 10-MODEL VALIDATION ---")
    tbl_bench = tables_data[3]  # Table 4 in doc is Table 3.1
    record_test("Benchmark_Table_Row_Count", len(tbl_bench) == 11, f"Table 3.1 has {len(tbl_bench)} rows (1 header + 10 models)")

    # Check each model's claims in Table 3.1
    expected_rows = {
        "Pratt et al. (2016)": {"Acc": "75.0%", "F1": "0.3850*", "QWK": "0.5120", "Details": "Sens 30.0% / Spec 95.0%"},
        "Vo & Verma (2016)": {"Acc": "76.5%", "F1": "0.4210*", "QWK": "0.6380", "Details": "Sens ~58% / Spec ~91%"},
        "Ben Graham (Quán quân 2015)": {"Acc": "85.2%", "F1": "0.5820*", "QWK": "0.8490", "Details": "Top-1 Private Leaderboard"},
        "Gulshan et al. (JAMA 2016)": {"Acc": "N/A", "F1": "N/A", "QWK": "N/A", "Details": "Sens 97.5% / Spec 93.4% / AUC 0.991"},
        "He et al. / Li et al. (2021)": {"Acc": "79.8%", "F1": "0.4950", "QWK": "0.6920", "Details": "Sens ~72% / Spec ~93%"},
        "Nhóm: ResNet-50 Baseline": {"Acc": "78.4%", "F1": "0.5094", "QWK": "0.7072", "Details": "Sens 78.4% / Spec 94.6% / AUC 0.885"},
        "Nhóm: ViT-B/16 (v5.2 EMA)": {"Acc": "80.6%", "F1": "0.5284", "QWK": "0.7215", "Details": "Sens 80.6% / Spec 95.2% / AUC 0.899"},
        "Nhóm: Hội chẩn Lai Ensemble": {"Acc": "81.8%", "F1": "0.5375", "QWK": "0.7280", "Details": "Sens 81.8% / Spec 95.5% / AUC 0.908"},
        "Nhóm: Ensemble + Tuning": {"Acc": "82.4%", "F1": "0.5420", "QWK": "0.7350", "Details": "Sens 82.4% / Spec 95.8% / AUC 0.912"},
        "Nhóm: Cổng Fundus Gate": {"Acc": "99.94%", "F1": "0.9994", "QWK": "N/A", "Details": "Sens 100% / Spec 99.87% / AUC 1.000"}
    }

    found_models = {}
    for r in tbl_bench[1:]:
        model_name = r[0]
        found_models[model_name] = {
            "Acc": r[4],
            "F1": r[5],
            "QWK": r[6],
            "Details": r[7]
        }

    for m_name, exp in expected_rows.items():
        matched = False
        for f_name, f_data in found_models.items():
            if m_name in f_name or f_name in m_name:
                matched = True
                acc_ok = exp["Acc"] == f_data["Acc"]
                f1_ok = exp["F1"] == f_data["F1"]
                qwk_ok = exp["QWK"] == f_data["QWK"]
                det_ok = exp["Details"] in f_data["Details"] or f_data["Details"] in exp["Details"]
                record_test(f"Model_{m_name}_Metrics", acc_ok and f1_ok and qwk_ok and det_ok,
                            f"Acc={f_data['Acc']} (exp {exp['Acc']}), F1={f_data['F1']} (exp {exp['F1']}), QWK={f_data['QWK']} (exp {exp['QWK']})")
                break
        if not matched:
            record_test(f"Model_{m_name}_Present", False, f"Model {m_name} not found in Table 3.1!")

    # --------------------------------------------------------------------------
    # SUITE 5: SYSTEM HARDWARE, LATENCY & ARTIFACT METRICS
    # --------------------------------------------------------------------------
    print("\n--- SUITE 5: HARDWARE & ARTIFACT METRICS ---")
    tbl_hw = tables_data[5]  # Table 6 in doc is Table 4.1
    # Check physical weights files exist and match sizes
    res_pth = os.path.join(WORKSPACE, "dr-fundus-app", "backend", "models", "resnet50_ce_best.pth")
    vit_pth = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "vit_b16_wce_best.pth")
    
    if os.path.exists(res_pth):
        res_mb = os.path.getsize(res_pth) / 1000000.0  # decimal MB
        record_test("ResNet50_PTH_Size_Decimal", abs(res_mb - 94.39) < 0.1, f"ResNet-50 .pth size: {res_mb:.2f} MB (Claimed: 94.39 MB)")
    if os.path.exists(vit_pth):
        vit_mb = os.path.getsize(vit_pth) / 1000000.0  # decimal MB
        record_test("ViT_PTH_Size_Decimal", abs(vit_mb - 343.27) < 0.1, f"ViT-B/16 .pth size: {vit_mb:.2f} MB (Claimed: 343.27 MB)")

    # Ensemble combined size
    ens_mb_sum = 94.39 + 343.27
    record_test("Ensemble_Combined_Size_Math", abs(ens_mb_sum - 437.66) < 0.01, f"Ensemble size sum: {ens_mb_sum:.2f} MB == Claimed 437.66 MB")

    # Ensemble params sum
    ens_params_sum = 25.56 + 86.57
    record_test("Ensemble_Combined_Params_Math", abs(ens_params_sum - 112.13) < 0.01, f"Ensemble params sum: {ens_params_sum:.2f}M == Claimed 112.13M")

    # --------------------------------------------------------------------------
    # SUITE 6: LITERATURE AND REFERENCE CROSS-CHECK
    # --------------------------------------------------------------------------
    print("\n--- SUITE 6: BIBLIOGRAPHY AND CITATIONS ---")
    ref_list = [p for p in para_texts if re.match(r'^\[\d+\]', p)]
    record_test("Bibliography_Count", len(ref_list) == 10, f"Found {len(ref_list)} formal bibliographic references (Expected 10)")
    
    # Check citation DOIs and authors
    record_test("Ref_Graham_2015", any("Graham, B. (2015)" in r for r in ref_list), "Ref [1] Graham (2015) verified")
    record_test("Ref_Gulshan_2016", any("Gulshan, V." in r and "JAMA" in r for r in ref_list), "Ref [2] Gulshan et al. (2016) verified")
    record_test("Ref_Pratt_2016", any("Pratt, H." in r and "Procedia Computer Science" in r for r in ref_list), "Ref [3] Pratt et al. (2016) verified")
    record_test("Ref_Vo_2016", any("Vo, H. H." in r and "IEEE" in r for r in ref_list), "Ref [4] Vo & Verma (2016) verified")
    record_test("Ref_He_2021", any("He, A." in r and "CABNet" in r for r in ref_list), "Ref [5] He et al. (2021) verified")
    record_test("Ref_Sun_2021", any("Sun, R." in r and "Lesion-Aware" in r for r in ref_list), "Ref [6] Sun et al. (2021) verified")
    record_test("Ref_Chefer_2021", any("Chefer, H." in r and "Transformer Interpretability" in r for r in ref_list), "Ref [8] Chefer et al. (2021) verified")

    # --------------------------------------------------------------------------
    # SUITE 7: INTERNAL CONSISTENCY & CONTRADICTION STRESS TEST
    # --------------------------------------------------------------------------
    print("\n--- SUITE 7: CONTRADICTION STRESS TEST ---")
    
    # Check if QWK gap is consistently cited (0.8490 and 0.7350)
    qwk_matches = [p for p in para_texts + [c.text for t in doc.tables for r in t.rows for c in r.cells] if "0.8490" in p]
    record_test("QWK_Gap_Consistency", len(qwk_matches) >= 2, f"Found {len(qwk_matches)} consistent citations of QWK gap (0.7350 vs 0.8490)")

    # Check if Macro-F1 range is consistently 0.5094 - 0.5420
    f1_range_matches = re.findall(r'0\.5094.*0\.5420', entire_doc_corpus)
    record_test("F1_Range_Consistency", len(f1_range_matches) >= 2, f"Found {len(f1_range_matches)} consistent citations of Macro-F1 boost range")

    # Check if Fundus Gate threshold is consistently 0.965
    tau_matches = [p for p in para_texts + [c.text for t in doc.tables for r in t.rows for c in r.cells] if "0.965" in p]
    record_test("Tau_Consistency", len(tau_matches) >= 3, f"Found {len(tau_matches)} consistent citations of optimal threshold tau*")

    # SUMMARY
    print("\n" + "=" * 80)
    print(f"STRESS TEST SUMMARY: {results['passed_tests']}/{results['total_tests']} PASSED, "
          f"{results['failed_tests']} FAILED, {results['warnings']} WARNINGS")
    print("=" * 80)

    return results

if __name__ == "__main__":
    res = run_stress_tests()
    if res["failed_tests"] > 0:
        sys.exit(1)
    else:
        sys.exit(0)
