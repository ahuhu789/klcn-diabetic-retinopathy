# -*- coding: utf-8 -*-
"""
EMPIRICAL CHALLENGER 3: INDEPENDENT COMPREHENSIVE VERIFICATION HARNESS
Target: Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx
Author: teamwork_preview_challenger_3
"""

import os
import sys
import zipfile
import xml.etree.ElementTree as ET
import docx

if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"c:\Users\thanh\Desktop\klcn\Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx"
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

report = {
    "passed": 0,
    "failed": 0,
    "warnings": 0,
    "checks": []
}

def record(test_id, passed, msg, warn=False):
    if passed:
        report["passed"] += 1
        status = "PASS"
    elif warn:
        report["warnings"] += 1
        status = "WARN"
    else:
        report["failed"] += 1
        status = "FAIL"
    report["checks"].append((test_id, status, msg))
    print(f"[{status}] {test_id}: {msg}")

def run_tests():
    print("=" * 80)
    print("EMPIRICAL CHALLENGER 3 - INDEPENDENT AUDIT HARNESS")
    print(f"Target file: {DOCX_PATH}")
    print("=" * 80)

    # ---------------------------------------------------------
    # SUITE A: File Integrity & XML Structure
    # ---------------------------------------------------------
    print("\n--- SUITE A: ARCHIVE INTEGRITY & XML VALIDATION ---")
    exists = os.path.exists(DOCX_PATH)
    record("A.1_File_Exists", exists, f"File existence at {DOCX_PATH}")
    if not exists:
        return report

    sz = os.path.getsize(DOCX_PATH)
    record("A.2_File_Size", sz >= 50000, f"File size: {sz} bytes ({sz/1024:.2f} KB, expected >= 50KB)")

    is_zip = zipfile.is_zipfile(DOCX_PATH)
    record("A.3_Zip_Valid", is_zip, "File is valid ZIP archive")
    if not is_zip:
        return report

    with zipfile.ZipFile(DOCX_PATH, 'r') as z:
        file_list = z.namelist()
        required_parts = ['word/document.xml', 'word/styles.xml', '[Content_Types].xml', '_rels/.rels', 'word/_rels/document.xml.rels']
        for part in required_parts:
            record(f"A.4_Part_{part}", part in file_list, f"OPC package contains {part}")

        doc_xml = z.read('word/document.xml')
        styles_xml = z.read('word/styles.xml')

    try:
        doc_dom = ET.fromstring(doc_xml)
        record("A.5_XML_WellFormed_Document", True, "word/document.xml parses cleanly as XML")
    except Exception as e:
        record("A.5_XML_WellFormed_Document", False, f"XML Parse error: {e}")
        return report

    try:
        styles_dom = ET.fromstring(styles_xml)
        record("A.6_XML_WellFormed_Styles", True, "word/styles.xml parses cleanly as XML")
    except Exception as e:
        record("A.6_XML_WellFormed_Styles", False, f"Styles Parse error: {e}")

    # ---------------------------------------------------------
    # SUITE B: Page Size & Margins (Standard A4)
    # ---------------------------------------------------------
    print("\n--- SUITE B: PAGE GEOMETRY & MARGINS (A4) ---")
    sect_elements = doc_dom.findall('.//w:sectPr', NS)
    record("B.1_SectPr_Count", len(sect_elements) == 1, f"Found {len(sect_elements)} w:sectPr element(s) (Expected 1)")

    sect = sect_elements[0] if sect_elements else None
    if sect is not None:
        pgSz = sect.find('w:pgSz', NS)
        record("B.2_PgSz_Exists", pgSz is not None, "w:pgSz tag exists in sectPr")
        if pgSz is not None:
            w_val = pgSz.attrib.get(W + 'w', '')
            h_val = pgSz.attrib.get(W + 'h', '')
            orient = pgSz.attrib.get(W + 'orient', 'portrait')
            record("B.3_PgSz_Width", w_val == "11906", f"w:w = '{w_val}' (Expected '11906' = 21.0cm)")
            record("B.4_PgSz_Height", h_val == "16838", f"w:h = '{h_val}' (Expected '16838' = 29.7cm)")
            record("B.5_PgSz_Orient", orient == "portrait", f"w:orient = '{orient}' (Expected 'portrait')")

        pgMar = sect.find('w:pgMar', NS)
        record("B.6_PgMar_Exists", pgMar is not None, "w:pgMar tag exists in sectPr")
        if pgMar is not None:
            top_m = pgMar.attrib.get(W + 'top', '')
            bottom_m = pgMar.attrib.get(W + 'bottom', '')
            left_m = pgMar.attrib.get(W + 'left', '')
            right_m = pgMar.attrib.get(W + 'right', '')
            # 3.5cm = 1984 dxa, 2.5cm = 1417 dxa
            record("B.7_Margin_Left", left_m == "1984", f"w:left = '{left_m}' dxa (Expected '1984' = 3.5cm)")
            record("B.8_Margin_Right", right_m == "1417", f"w:right = '{right_m}' dxa (Expected '1417' = 2.5cm)")
            record("B.9_Margin_Top", top_m == "1417", f"w:top = '{top_m}' dxa (Expected '1417' = 2.5cm)")
            record("B.10_Margin_Bottom", bottom_m == "1417", f"w:bottom = '{bottom_m}' dxa (Expected '1417' = 2.5cm)")

    # ---------------------------------------------------------
    # SUITE C: Table Cell Widths (Zero Duplicate tcW)
    # ---------------------------------------------------------
    print("\n--- SUITE C: CELL WIDTH CONFORMANCE & DUPLICATE TCW ---")
    all_cells = doc_dom.findall('.//w:tc', NS)
    record("C.1_Total_Cells_Count", len(all_cells) == 229, f"Total table cells = {len(all_cells)} (Expected 229)")

    duplicate_tcw_cells = []
    missing_tcw_cells = []
    single_tcw_cells = 0

    for idx, c in enumerate(all_cells):
        tcPr = c.find('w:tcPr', NS)
        if tcPr is None:
            missing_tcw_cells.append(idx)
            continue
        tcw_list = tcPr.findall('w:tcW', NS)
        if len(tcw_list) > 1:
            duplicate_tcw_cells.append((idx, len(tcw_list)))
        elif len(tcw_list) == 1:
            single_tcw_cells += 1
        else:
            missing_tcw_cells.append(idx)

    record("C.2_Zero_Duplicate_tcW", len(duplicate_tcw_cells) == 0,
           f"Duplicate tcW cells count: {len(duplicate_tcw_cells)} (Violations: {duplicate_tcw_cells})")
    record("C.3_Zero_Missing_tcW", len(missing_tcw_cells) == 0,
           f"Missing tcW cells count: {len(missing_tcw_cells)} (Missing: {missing_tcw_cells})")
    record("C.4_Single_tcW_Per_Cell", single_tcw_cells == 229,
           f"Cells with exactly 1 valid tcW: {single_tcw_cells} / 229")

    # ---------------------------------------------------------
    # SUITE D: Table Grid & Synchronization Across All 7 Tables
    # ---------------------------------------------------------
    print("\n--- SUITE D: TBLGRID & ROW WIDTH SYNCHRONIZATION ---")
    tables = doc_dom.findall('.//w:tbl', NS)
    record("D.1_Table_Count", len(tables) == 7, f"Found {len(tables)} tables (Expected 7)")

    table_cell_counts = [10, 66, 32, 88, 1, 30, 2]

    for t_idx, t in enumerate(tables):
        t_num = t_idx + 1
        tblGrid = t.find('w:tblGrid', NS)
        record(f"D.2_T{t_num}_TblGrid_Exists", tblGrid is not None, f"Table {t_num} has w:tblGrid")
        if tblGrid is None:
            continue

        gridCols = tblGrid.findall('w:gridCol', NS)
        gridCol_widths = [int(gc.attrib.get(W + 'w', 0)) for gc in gridCols]
        grid_sum = sum(gridCol_widths)
        
        # 15.00cm is ~8502-8504 twips
        record(f"D.3_T{t_num}_GridSum", 8502 <= grid_sum <= 8504,
               f"Table {t_num} gridCol sum = {grid_sum} twips ({grid_sum/566.929:.3f} cm, expected 8502-8504)")

        # Verify EVERY row in this table matches gridCol widths!
        rows = t.findall('w:tr', NS)
        all_rows_synced = True
        mismatched_rows = []

        total_t_cells = sum(len(r.findall('w:tc', NS)) for r in rows)
        exp_cells = table_cell_counts[t_idx]
        record(f"D.4_T{t_num}_CellCount", total_t_cells == exp_cells,
               f"Table {t_num} cell count = {total_t_cells} (Expected {exp_cells})")

        for r_idx, r in enumerate(rows):
            row_cells = r.findall('w:tc', NS)
            row_tcw = []
            for c in row_cells:
                tcW = c.find('w:tcPr/w:tcW', NS)
                if tcW is not None:
                    row_tcw.append(int(tcW.attrib.get(W + 'w', 0)))
                else:
                    row_tcw.append(-1)

            if row_tcw != gridCol_widths:
                all_rows_synced = False
                mismatched_rows.append((r_idx, row_tcw, gridCol_widths))

        record(f"D.5_T{t_num}_All_Rows_Synced", all_rows_synced,
               f"Table {t_num} ({len(rows)} rows): All rows match tblGrid exactly. Mismatches: {len(mismatched_rows)}")

    # ---------------------------------------------------------
    # SUITE E: Table Layout Properties (cantSplit, tblHeader, shd, borders, margins)
    # ---------------------------------------------------------
    print("\n--- SUITE E: ADVANCED OOXML TABLE LAYOUT PROPERTIES ---")
    for t_idx, t in enumerate(tables):
        t_num = t_idx + 1
        rows = t.findall('w:tr', NS)
        
        # Check cantSplit on every row
        rows_without_cantsplit = [r_idx for r_idx, r in enumerate(rows) if r.find('w:trPr/w:cantSplit', NS) is None]
        record(f"E.1_T{t_num}_CantSplit_AllRows", len(rows_without_cantsplit) == 0,
               f"Table {t_num}: cantSplit on all {len(rows)} rows. Missing rows: {rows_without_cantsplit}")

        # Check tblHeader on row 0
        r0_trPr = rows[0].find('w:trPr', NS) if rows else None
        has_tbl_header = r0_trPr is not None and r0_trPr.find('w:tblHeader', NS) is not None
        record(f"E.2_T{t_num}_TblHeader", has_tbl_header,
               f"Table {t_num} row 0 has w:tblHeader for repeated header: {has_tbl_header}")

        # Check shading (shd)
        cells_with_shd = sum(1 for c in t.findall('.//w:tc', NS) if c.find('w:tcPr/w:shd', NS) is not None)
        record(f"E.3_T{t_num}_Shading", cells_with_shd > 0,
               f"Table {t_num} has styling shading: {cells_with_shd} shaded cells")

        # Check cell borders (tcBorders)
        cells_with_borders = sum(1 for c in t.findall('.//w:tc', NS) if c.find('w:tcPr/w:tcBorders', NS) is not None)
        record(f"E.4_T{t_num}_Borders", cells_with_borders > 0,
               f"Table {t_num} has custom borders: {cells_with_borders} bordered cells")

        # Check cell margins (tcMar)
        cells_with_mar = sum(1 for c in t.findall('.//w:tc', NS) if c.find('w:tcPr/w:tcMar', NS) is not None)
        record(f"E.5_T{t_num}_CellPadding", cells_with_mar > 0,
               f"Table {t_num} has custom padding (tcMar): {cells_with_mar} padded cells")

    # ---------------------------------------------------------
    # SUITE F: Typography & Font Inheritance
    # ---------------------------------------------------------
    print("\n--- SUITE F: TYPOGRAPHY AUDIT & RUN-LEVEL FONT INSPECTION ---")
    doc_docx = docx.Document(DOCX_PATH)

    # Style Normal font in python-docx API
    normal_font = doc_docx.styles['Normal'].font.name
    record("F.1_Style_Normal_Font", normal_font == "Times New Roman",
           f"Style 'Normal' default font: '{normal_font}'")

    # Inspect Normal style in styles.xml
    normal_style_el = None
    for st in styles_dom.findall('.//w:style', NS):
        if st.attrib.get(W + 'styleId') == 'Normal':
            normal_style_el = st
            break
    
    normal_rf = normal_style_el.find('.//w:rFonts', NS) if normal_style_el is not None else None
    normal_ascii = normal_rf.attrib.get(W + 'ascii', '') if normal_rf is not None else ''
    normal_hansi = normal_rf.attrib.get(W + 'hAnsi', '') if normal_rf is not None else ''
    record("F.2_Styles_XML_Normal_TNR", normal_ascii == "Times New Roman" and normal_hansi == "Times New Roman",
           f"styles.xml Normal style defines ascii='{normal_ascii}', hAnsi='{normal_hansi}'")

    # Inspect all runs in document.xml
    all_runs = doc_dom.findall('.//w:r', NS)
    record("F.3_Runs_Present", len(all_runs) > 0, f"Found {len(all_runs)} text runs in document.xml")

    runs_with_explicit_font = 0
    runs_with_non_tnr_font = []
    runs_with_no_rfonts = 0

    for idx, r in enumerate(all_runs):
        rFonts = r.find('w:rPr/w:rFonts', NS)
        if rFonts is not None:
            ascii_font = rFonts.attrib.get(W + 'ascii', '')
            hAnsi_font = rFonts.attrib.get(W + 'hAnsi', '')
            cs_font = rFonts.attrib.get(W + 'cs', '')
            if any(f and "Times New Roman" not in f for f in [ascii_font, hAnsi_font, cs_font]):
                runs_with_non_tnr_font.append((idx, ascii_font, hAnsi_font, cs_font))
            else:
                runs_with_explicit_font += 1
        else:
            runs_with_no_rfonts += 1

    record("F.4_Explicit_Runs_TNR", len(runs_with_non_tnr_font) == 0,
           f"Zero runs with non-Times New Roman fonts. Violations: {runs_with_non_tnr_font}")
    record("F.5_Run_Coverage", runs_with_explicit_font == len(all_runs),
           f"100% of runs ({runs_with_explicit_font}/{len(all_runs)}) explicitly specify Times New Roman")

    # ---------------------------------------------------------
    # SUITE G: Text Corpus & Core Requirements (R1, R2, R3, R4)
    # ---------------------------------------------------------
    print("\n--- SUITE G: SCIENTIFIC AND REGULATORY CONTENT VERIFICATION ---")
    doc_text_parts = [p.text for p in doc_docx.paragraphs] + [c.text for t in doc_docx.tables for r in t.rows for c in r.cells]
    doc_corpus = "\n".join(doc_text_parts)

    req_checks = [
        ("G.1_R1_Ben_Graham_2015", "Ben Graham (Quán quân 2015)" in doc_corpus and "0.8490" in doc_corpus, "Ben Graham 2015 Kaggle Winner QWK 0.8490"),
        ("G.2_R1_Gulshan_JAMA_2016", "Gulshan et al. (JAMA 2016)" in doc_corpus and "0.991" in doc_corpus, "Gulshan et al. 2016 AUC 0.991"),
        ("G.3_R1_Pratt_2016", "Pratt et al. (2016)" in doc_corpus and "0.5120" in doc_corpus and "30.0%" in doc_corpus, "Pratt 2016 Acc 75.0%, QWK 0.5120, Sens 30.0%"),
        ("G.4_R1_Vo_Verma_2016", "Vo & Verma (2016)" in doc_corpus and "0.6380" in doc_corpus and "76.5%" in doc_corpus, "Vo & Verma 2016 Acc 76.5%, QWK 0.6380"),
        ("G.5_R1_Modern_Baselines", "He et al. / Li et al. (2021)" in doc_corpus and "0.6920" in doc_corpus, "He / Li 2021 QWK 0.6920"),
        ("G.6_R1_ResNet50_Baseline", "Nhóm: ResNet-50 Baseline" in doc_corpus and "0.7072" in doc_corpus and "0.5094" in doc_corpus, "ResNet-50 QWK 0.7072, F1 0.5094"),
        ("G.7_R1_ViT_B16_v5.2", "Nhóm: ViT-B/16 (v5.2 EMA)" in doc_corpus and "0.7215" in doc_corpus and "0.5284" in doc_corpus, "ViT-B/16 QWK 0.7215, F1 0.5284"),
        ("G.8_R1_Ensemble_Hybrid", "Nhóm: Hội chẩn Lai Ensemble" in doc_corpus and "0.7280" in doc_corpus, "Ensemble QWK 0.7280"),
        ("G.9_R1_Ensemble_Tuned", "Nhóm: Ensemble + Tuning" in doc_corpus and "0.7350" in doc_corpus and "0.5420" in doc_corpus, "Ensemble + Tuning QWK 0.7350, F1 0.5420"),
        ("G.10_R1_Fundus_Gate", "Nhóm: Cổng Fundus Gate" in doc_corpus and "99.94%" in doc_corpus and "100%" in doc_corpus, "Fundus Gate 99.94% Acc, 100% Sens"),
        ("G.11_R2_Loss_Smoothing", "Smoothed Weighted Cross-Entropy Loss" in doc_corpus and ("w_c = \\sqrt{w_c^{raw}}" in doc_corpus or "căn bậc hai" in doc_corpus), "Smoothed WCE Loss mathematical explanation"),
        ("G.12_R2_Dual_XAI", "Grad-CAM" in doc_corpus and "Transformer Attribution" in doc_corpus, "Dual XAI interpretability"),
        ("G.13_R2_Fundus_Gate_Specs", "99.87%" in doc_corpus and "0.11%" in doc_corpus and "0.965" in doc_corpus, "Fundus gate Spec 99.87%, Challenge FPR 0.11%, tau* 0.965"),
        ("G.14_R2_FullStack_Deployment", "React" in doc_corpus and "FastAPI" in doc_corpus and "Cloudflare" in doc_corpus, "Fullstack Web application and Cloudflare Tunnel"),
        ("G.15_R3_QWK_Gap_Analysis", "0.8490" in doc_corpus and "Dual" in doc_corpus and "Tesla T4" in doc_corpus, "QWK gap analysis vs Graham 2015 infrastructure"),
        ("G.16_R3_Grade0_vs_1", "Độ 0" in doc_corpus and "Độ 1" in doc_corpus and "vi phình mạch" in doc_corpus, "Grade 0 vs 1 microaneurysm resolution analysis"),
        ("G.17_R3_ViT_Latency_Cost", "343" in doc_corpus and "CPU" in doc_corpus and "520" in doc_corpus, "ViT 343MB model size and CPU inference latency"),
        ("G.18_R3_External_Validation", "Messidor-2" in doc_corpus and "APTOS 2019" in doc_corpus, "External dataset validation discussion"),
        ("G.19_R4_HUIT_Headers", "BỘ CÔNG THƯƠNG" in doc_corpus and "TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH" in doc_corpus and "KHOA CÔNG NGHỆ THÔNG TIN" in doc_corpus, "HUIT formal institutional headers"),
        ("G.20_R4_Advisor_And_Team", "ThS. Trần Đình Toàn" in doc_corpus and "Nguyễn Huy Thanh" in doc_corpus and "Nguyễn Cao Tấn Phúc" in doc_corpus and "Nguyễn Chí Thành" in doc_corpus, "Advisor and student names"),
        ("G.21_R4_Signature_Block", "XÁC NHẬN CỦA GIẢNG VIÊN HƯỚNG DẪN" in doc_corpus and "ĐẠI DIỆN NHÓM SINH VIÊN THỰC HIỆN" in doc_corpus, "Formal 2-column signature block")
    ]

    for c_id, cond, msg in req_checks:
        record(c_id, cond, msg)

    print("\n" + "=" * 80)
    print(f"FINAL AUDIT SUMMARY: {report['passed']} PASSED, {report['failed']} FAILED, {report['warnings']} WARNINGS")
    print("=" * 80)
    return report

if __name__ == '__main__':
    rep = run_tests()
    if rep['failed'] > 0:
        sys.exit(1)
    else:
        sys.exit(0)
