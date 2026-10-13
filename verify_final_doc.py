import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

final_path = r"c:\Users\thanh\Desktop\klcn\Bao_Cao_Thuc_Tap_Nghe_Nghiep_Nguyen_Huy_Thanh_DEMO_FINAL.docx"
doc = docx.Document(final_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- CHECKING DANH MỤC HÌNH ---")
dm_start = False
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt == "DANH MỤC HÌNH":
        dm_start = True
        print(f"P[{i}]: {txt}")
        continue
    if dm_start:
        if "DANH MỤC BẢNG" in txt:
            print(f"P[{i}]: {txt}")
            break
        if txt.startswith("Hình"):
            print(f"  {txt}")

print("\n--- CHECKING CAPTIONS IN BODY TEXT ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Hình") and i > 100:
        # Check image paragraph
        prev_p = doc.paragraphs[i-1]
        blips = prev_p._p.xpath('.//a:blip')
        extents = prev_p._p.xpath('.//wp:extent')
        dim_str = ""
        if extents:
            cx = int(extents[0].get('cx'))
            cy = int(extents[0].get('cy'))
            dim_str = f"({cx/360000:.2f} cm x {cy/360000:.2f} cm)"
        print(f"P[{i}] (img_blips={len(blips)} {dim_str}): {txt}")

print("\n--- CHECKING INTEGRITY OF CRITICAL FIELDS ---")
# Student info
for p in doc.paragraphs[:40]:
    if any(k in p.text for k in ["Nguyễn Huy Thanh", "2001230885", "Tâm Anh", "Nguyễn Văn Trắng"]):
        print(f"Confirmed: {p.text.strip()}")

print("\nVerification complete!")
