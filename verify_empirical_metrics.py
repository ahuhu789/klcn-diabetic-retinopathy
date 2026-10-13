# -*- coding: utf-8 -*-
"""
EMPIRICAL CHALLENGER VERIFICATION SCRIPT
Target: Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx
Purpose: Extract all quantitative claims and tables, cross-check against
         codebase files, raw notebooks, JSON/CSV metrics, and literature.
"""

import sys
import os
import json
import re
import math
import docx
import pandas as pd

if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"c:\Users\thanh\Desktop\klcn\Bao_Cao_Tien_Do_Tuan_KLCN_28-09_den_05-10.docx"
WORKSPACE = r"c:\Users\thanh\Desktop\klcn"

def main():
    print("=" * 80)
    print("EMPIRICAL STRESS-TESTING OF QUANTITATIVE METRICS")
    print("=" * 80)

    # 1. Parse Word Document
    doc = docx.Document(DOCX_PATH)
    print(f"Loaded document: {DOCX_PATH}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}, Total Tables: {len(doc.tables)}")

    # Extract all text
    all_text = []
    for idx, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            all_text.append((f"P{idx}", t))

    table_data = []
    for t_idx, tbl in enumerate(doc.tables):
        rows = []
        for r_idx, row in enumerate(tbl.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            rows.append(cells)
        table_data.append(rows)

    # Print tables summary
    print("\n--- EXTRACTED TABLES SUMMARY ---")
    for t_idx, rows in enumerate(table_data):
        print(f"Table {t_idx+1}: {len(rows)} rows x {len(rows[0]) if rows else 0} cols")
        for r_idx, r in enumerate(rows[:3]):  # show first 3 rows
            print(f"  R{r_idx}: {' | '.join(r[:5])}")
        if len(rows) > 3:
            print(f"  ... ({len(rows)-3} more rows)")

    # 2. Source of Truth Checks
    print("\n" + "=" * 80)
    print("CHECK 1: FUNDUS GATE METRICS (metrics.json vs docx)")
    print("=" * 80)
    
    fg_json_path = os.path.join(WORKSPACE, "fundus-gate", "metrics.json")
    if not os.path.exists(fg_json_path):
        fg_json_path = os.path.join(WORKSPACE, "fundus-gate", "output", "metrics.json")
    
    with open(fg_json_path, "r", encoding="utf-8") as f:
        fg_data = json.load(f)

    # Ground truth values from JSON
    fg_gt_recall = fg_data["test_metrics"]["recall"]
    fg_gt_spec = fg_data["test_metrics"]["specificity"]
    fg_gt_acc = fg_data["test_metrics"]["accuracy"]
    fg_gt_fpr = fg_data["test_metrics"]["fpr"]
    fg_gt_f1 = fg_data["test_metrics"]["f1"]
    fg_gt_tau = fg_data["optimal_threshold"]
    fg_gt_params = fg_data["total_parameters"]
    fg_gt_cm = fg_data["test_metrics"]["confusion_matrix"]
    
    ch_gt_total = fg_data["challenge_test_metrics"]["total_samples"]
    ch_gt_tn = fg_data["challenge_test_metrics"]["true_negatives"]
    ch_gt_fp = fg_data["challenge_test_metrics"]["false_positives"]
    ch_gt_spec = fg_data["challenge_test_metrics"]["challenge_specificity"]
    ch_gt_fpr = fg_data["challenge_test_metrics"]["challenge_fpr"]
    
    art_keras = fg_data["artifacts"]["keras_size_mb"]
    art_tflite = fg_data["artifacts"]["tflite_size_mb"]
    art_fp16 = fg_data["artifacts"]["tflite_fp16_size_mb"]

    print("Ground Truth Fundus Gate JSON:")
    print(f"  Recall: {fg_gt_recall * 100:.2f}%")
    print(f"  Specificity: {fg_gt_spec * 100:.2f}%")
    print(f"  Accuracy: {fg_gt_acc * 100:.2f}%")
    print(f"  FPR: {fg_gt_fpr * 100:.2f}%")
    print(f"  F1: {fg_gt_f1:.4f}")
    print(f"  Optimal Threshold tau*: {fg_gt_tau}")
    print(f"  Total Params: {fg_gt_params}")
    print(f"  CM: {fg_gt_cm}")
    print(f"  Challenge Spec: {ch_gt_spec * 100:.2f}%, FPR: {ch_gt_fpr * 100:.2f}% (TN={ch_gt_tn}, FP={ch_gt_fp}, Total={ch_gt_total})")
    print(f"  Artifact sizes: Keras={art_keras}MB, TFLite={art_tflite}MB, FP16={art_fp16}MB")

    # Math recalculation of Fundus Gate metrics:
    tp, fp, tn, fn = fg_gt_cm["tp"], fg_gt_cm["fp"], fg_gt_cm["tn"], fg_gt_cm["fn"]
    calc_rec = tp / (tp + fn)
    calc_spec = tn / (tn + fp)
    calc_acc = (tp + tn) / (tp + tn + fp + fn)
    calc_fpr = fp / (fp + tn)
    calc_prec = tp / (tp + fp)
    calc_f1 = 2 * (calc_prec * calc_rec) / (calc_prec + calc_rec)
    calc_ch_spec = ch_gt_tn / ch_gt_total
    calc_ch_fpr = ch_gt_fp / ch_gt_total

    print("\nMathematical Re-computation of Confusion Matrix:")
    print(f"  Calculated Recall: {calc_rec:.6f} == {fg_gt_recall:.6f} -> {'PASS' if abs(calc_rec - fg_gt_recall) < 1e-5 else 'FAIL'}")
    print(f"  Calculated Specificity: {calc_spec:.6f} == {fg_gt_spec:.6f} -> {'PASS' if abs(calc_spec - fg_gt_spec) < 1e-4 else 'FAIL'}")
    print(f"  Calculated Accuracy: {calc_acc:.6f} == {fg_gt_acc:.6f} -> {'PASS' if abs(calc_acc - fg_gt_acc) < 1e-4 else 'FAIL'}")
    print(f"  Calculated FPR: {calc_fpr:.6f} == {fg_gt_fpr:.6f} -> {'PASS' if abs(calc_fpr - fg_gt_fpr) < 1e-4 else 'FAIL'}")
    print(f"  Calculated F1: {calc_f1:.6f} == {fg_gt_f1:.6f} -> {'PASS' if abs(calc_f1 - fg_gt_f1) < 1e-4 else 'FAIL'}")
    print(f"  Calculated Challenge Spec: {calc_ch_spec:.6f} == {ch_gt_spec:.6f} -> {'PASS' if abs(calc_ch_spec - ch_gt_spec) < 1e-4 else 'FAIL'}")
    print(f"  Calculated Challenge FPR: {calc_ch_fpr:.6f} == {ch_gt_fpr:.6f} -> {'PASS' if abs(calc_ch_fpr - ch_gt_fpr) < 1e-4 else 'FAIL'}")

    # Physical file sizes
    fp16_path = os.path.join(WORKSPACE, "fundus-gate", "fundus_gate_float16.tflite")
    if os.path.exists(fp16_path):
        fp16_bytes = os.path.getsize(fp16_path)
        fp16_mb = fp16_bytes / (1024 * 1024)
        print(f"  Physical FP16 TFLite file size: {fp16_bytes} bytes = {fp16_mb:.2f} MB (Expected ~1.82 MB)")

    print("\n" + "=" * 80)
    print("CHECK 2: CLASSIFICATION METRICS (results_comparison.csv & metrics.py vs docx)")
    print("=" * 80)
    
    csv_path = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "results_comparison.csv")
    df = pd.read_csv(csv_path)
    print("CSV contents:")
    print(df.to_string())

    # Check best metrics in CSV
    best_res_loss = df["ResNet50_Loss"].min()
    best_res_f1 = df["ResNet50_F1"].max()
    best_res_qwk = df["ResNet50_QWK"].max()
    
    # Identify epochs
    epoch_best_loss = df.loc[df["ResNet50_Loss"] == best_res_loss, "Epoch"].values[0]
    epoch_best_f1 = df.loc[df["ResNet50_F1"] == best_res_f1, "Epoch"].values[0]
    epoch_best_qwk = df.loc[df["ResNet50_QWK"] == best_res_qwk, "Epoch"].values[0]

    print(f"\nResNet-50 in CSV:")
    print(f"  Best Val Loss: {best_res_loss:.4f} at Epoch {epoch_best_loss}")
    print(f"  Best Macro-F1: {best_res_f1:.4f} at Epoch {epoch_best_f1}")
    print(f"  Best QWK: {best_res_qwk:.4f} at Epoch {epoch_best_qwk}")

    # Check metrics.py
    metrics_py_path = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "metrics.py")
    with open(metrics_py_path, "r", encoding="utf-8") as f:
        metrics_py_content = f.read()

    print("\nmetrics.py ground truth:")
    for line in metrics_py_content.splitlines():
        if any(k in line for k in ["Best Macro F-1", "Best QWK", "Best Val Loss", "Do chinh xac", "QWK cao nhat", "ViT-B/16"]):
            print(f"  {line.strip()}")

    # Physical model sizes
    res_pth = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "resnet50_ce_best.pth")
    vit_pth = os.path.join(WORKSPACE, "eye-dr-gradio-demo", "vit_b16_wce_best.pth")
    if os.path.exists(res_pth):
        res_bytes = os.path.getsize(res_pth)
        print(f"  ResNet-50 physical .pth size: {res_bytes} bytes = {res_bytes/(1024*1024):.2f} MB")
    if os.path.exists(vit_pth):
        vit_bytes = os.path.getsize(vit_pth)
        print(f"  ViT-B/16 physical .pth size: {vit_bytes} bytes = {vit_bytes/(1024*1024):.2f} MB")

    print("\n" + "=" * 80)
    print("CHECK 3: KAGGLE NOTEBOOK EVALUATION LOGS & LOSS WEIGHTS")
    print("=" * 80)
    
    nb_path = os.path.join(WORKSPACE, "kaggle_notebook_pull", "klcn-resnet50-baseline-part2.ipynb")
    with open(nb_path, "r", encoding="utf-8") as f:
        nb_data = json.load(f)

    # Class distribution
    # Class counts: [25810, 2443, 5292, 873, 708]
    class_counts = [25810, 2443, 5292, 873, 708]
    total_samples = sum(class_counts)
    num_classes = len(class_counts)

    raw_weights = [total_samples / (num_classes * c) for c in class_counts]
    smoothed_weights = [math.sqrt(w) for w in raw_weights]

    print(f"Class distribution: {class_counts}, Total: {total_samples}")
    print(f"Raw weights: {[round(w, 4) for w in raw_weights]}")
    print(f"  Raw max/min ratio: {max(raw_weights) / min(raw_weights):.2f}x ({max(raw_weights):.3f} vs {min(raw_weights):.3f})")
    print(f"Smoothed weights: {[round(w, 4) for w in smoothed_weights]}")
    print(f"  Smoothed max/min ratio: {max(smoothed_weights) / min(smoothed_weights):.2f}x ({max(smoothed_weights):.3f} vs {min(smoothed_weights):.3f})")

    # Inspect cells for metrics / outputs
    print("\nSearching notebook outputs for validation metrics:")
    found_metrics_outputs = []
    for cell_idx, cell in enumerate(nb_data.get("cells", [])):
        cell_type = cell.get("cell_type")
        source = "".join(cell.get("source", []))
        outputs = cell.get("outputs", [])
        
        # Check source keywords
        if any(kw in source for kw in ["OptimizedRounder", "quadratic_weighted_kappa", "classification_report", "ModelEMA"]):
            print(f"  [Cell {cell_idx} ({cell_type})] contains key model/eval code.")
            for out in outputs:
                if "text" in out:
                    out_text = "".join(out["text"])
                    for line in out_text.splitlines():
                        if any(k in line for k in ["QWK", "Kappa", "Macro", "Accuracy", "Epoch", "Thresholds"]):
                            found_metrics_outputs.append(f"Cell {cell_idx}: {line.strip()}")
                elif "data" in out and "text/plain" in out["data"]:
                    out_text = "".join(out["data"]["text/plain"])
                    for line in out_text.splitlines():
                        if any(k in line for k in ["QWK", "Kappa", "Macro", "Accuracy"]):
                            found_metrics_outputs.append(f"Cell {cell_idx}: {line.strip()}")

    print(f"  Found {len(found_metrics_outputs)} matching output lines.")
    for l in found_metrics_outputs[:10]:
        print(f"    {l}")

    print("\n" + "=" * 80)
    print("CHECK 4: LITERATURE CITATIONS & BENCHMARKS CROSS-CHECK")
    print("=" * 80)
    lit_survey_path = os.path.join(WORKSPACE, ".agents", "teamwork", "teamwork_preview_explorer_survey_2", "survey_literature_benchmarks.md")
    with open(lit_survey_path, "r", encoding="utf-8") as f:
        lit_content = f.read()

    expected_lit = {
        "Ben Graham (2015)": {"QWK": 0.8490, "Model": "SparseConvNet"},
        "Gulshan et al. (2016)": {"AUC": 0.991, "High Sens": 0.975, "High Spec": 0.934},
        "Pratt et al. (2016)": {"Acc": 0.750, "Sens": 0.300, "QWK": 0.5120, "Macro-F1": 0.3850},
        "Vo & Verma (2016)": {"Acc": 0.765, "QWK": 0.6380, "Macro-F1": 0.4210},
        "He et al. (2021)": {"Acc": 0.798, "QWK": 0.6920, "Macro-F1": 0.4950}
    }

    for author, metrics in expected_lit.items():
        print(f"Checking literature ground truth for {author}:")
        for m_name, m_val in metrics.items():
            if str(m_val) in lit_content:
                print(f"  - {m_name}: {m_val} -> CONFIRMED in literature survey")
            else:
                print(f"  - {m_name}: {m_val} -> NOT FOUND in literature survey!")

    # Check presence in docx
    full_doc_text = " ".join([t for _, t in all_text])
    for rows in table_data:
        for r in rows:
            full_doc_text += " " + " ".join(r)

    print("\nVerifying literature figures in DOCX:")
    for author, metrics in expected_lit.items():
        auth_key = author.split()[0]
        in_docx = auth_key in full_doc_text
        print(f"  {author} present in docx: {in_docx}")
        for m_name, m_val in metrics.items():
            val_str = str(m_val)
            # check percentage or decimal
            if isinstance(m_val, float) and m_val < 1.0:
                pct_str = f"{m_val*100:.1f}%"
            else:
                pct_str = "N/A"
            found_val = (val_str in full_doc_text) or (pct_str in full_doc_text)
            print(f"    - {m_name} ({val_str} / {pct_str}): {'FOUND' if found_val else 'MISSING'}")

    print("\n" + "=" * 80)
    print("CHECK 5: SYSTEMATIC CROSS-REFERENCE OF ALL METRIC CLAIMS IN DOCX")
    print("=" * 80)

    # Let's check Table 3.1 verbatim from the document
    tbl_3_1 = None
    for tbl in table_data:
        if any("Pratt et al." in " ".join(r) for r in tbl):
            tbl_3_1 = tbl
            break

    if tbl_3_1:
        print("Table 3.1 found with", len(tbl_3_1), "rows:")
        headers = tbl_3_1[0]
        print("  Headers:", headers)
        for r in tbl_3_1[1:]:
            print("  Row:", r[0], "-->", " | ".join(r[1:]))
    else:
        print("ERROR: Table 3.1 not found!")

if __name__ == "__main__":
    main()
